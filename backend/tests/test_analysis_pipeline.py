import unittest
import uuid
import io
import docx
from fastapi.testclient import TestClient
from app.main import app
from app.core.database import Base, engine

class TestAnalysisPipeline(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.client = TestClient(app)
        Base.metadata.create_all(bind=engine)

    def test_seeker_analysis_pipeline(self):
        """Tests resume upload, analysis processing, and result retrieval."""
        user_email = f"pipeline_{uuid.uuid4().hex[:8]}@example.com"
        reg = self.client.post("/api/v1/auth/register", json={
            "name": "Pipeline Tester",
            "email": user_email,
            "password": "Password123!",
            "role": "JOB_SEEKER"
        })
        self.assertEqual(reg.status_code, 201)

        login = self.client.post("/api/v1/auth/login", json={
            "email": user_email,
            "password": "Password123!"
        })
        self.assertEqual(login.status_code, 200)
        token = login.json()["access_token"]
        headers = {"Authorization": f"Bearer {token}"}

        # 1. Create a genuine DOCX resume
        doc = docx.Document()
        doc.add_heading("Jane Doe", level=1)
        doc.add_paragraph("Software Engineer with 4 years of experience in Python, FastAPI, and Docker.")
        doc.add_paragraph("Key Qualifications: Python, FastAPI, Docker, PostgreSQL.")
        doc.add_paragraph("Experience: Senior Backend Developer at Acme Systems building scalable microservices in Python.")

        docx_buf = io.BytesIO()
        doc.save(docx_buf)
        docx_buf.seek(0)

        files = {
            "file": ("jane_doe_resume.docx", docx_buf, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
        }
        res_upload = self.client.post("/api/v1/resumes/upload", headers=headers, files=files)
        self.assertEqual(res_upload.status_code, 201)
        resume_id = res_upload.json()["id"]

        # 2. Run Fit Analysis with custom Job Description
        jd_text = (
            "Role: Backend Engineer\n"
            "Requirements:\n"
            "- 3+ years experience with Python and FastAPI\n"
            "- Experience with PostgreSQL\n"
            "Preferred:\n"
            "- Kubernetes"
        )
        analysis_payload = {
            "resume_id": resume_id,
            "job_description": jd_text,
            "job_title": "Backend Engineer",
            "company_name": "Acme Corp"
        }
        res_analysis = self.client.post("/api/v1/analyses/candidate", headers=headers, json=analysis_payload)
        self.assertEqual(res_analysis.status_code, 201)
        analysis_data = res_analysis.json()
        analysis_id = analysis_data["id"]

        # 3. Retrieve analysis detail
        res_detail = self.client.get(f"/api/v1/analyses/{analysis_id}", headers=headers)
        self.assertEqual(res_detail.status_code, 200)
        detail_data = res_detail.json()

        self.assertIn(detail_data["status"], ["COMPLETED", "ANALYZING", "PROCESSING"])
        if detail_data["status"] == "COMPLETED":
            result = detail_data["match_result"]
            self.assertIsNotNone(result)
            self.assertIn("component_scores", result)
            self.assertGreater(result["overall_score"], 0)

if __name__ == "__main__":
    unittest.main()
