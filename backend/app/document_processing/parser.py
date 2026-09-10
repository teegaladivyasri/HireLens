import io
from abc import ABC, abstractmethod
from typing import Optional
import pypdf
import docx


class DocumentParser(ABC):
    """Abstract base class for document text parsers."""

    @abstractmethod
    def parse(self, file_bytes: bytes) -> str:
        """Parse raw file bytes into extracted plain text."""
        pass


class PDFParser(DocumentParser):
    """PDF document parser using pypdf."""

    def parse(self, file_bytes: bytes) -> str:
        text_parts = []
        try:
            reader = pypdf.PdfReader(io.BytesIO(file_bytes))
            for page in reader.pages:
                page_text = page.extract_text()
                if page_text:
                    text_parts.append(page_text)
        except Exception as e:
            raise ValueError(f"Failed to parse PDF document: {str(e)}")
        return "\n".join(text_parts).strip()


class DOCXParser(DocumentParser):
    """DOCX document parser using python-docx."""

    def parse(self, file_bytes: bytes) -> str:
        text_parts = []
        try:
            doc = docx.Document(io.BytesIO(file_bytes))
            for paragraph in doc.paragraphs:
                if paragraph.text.strip():
                    text_parts.append(paragraph.text.strip())
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                    if row_text:
                        text_parts.append(row_text)
        except Exception as e:
            raise ValueError(f"Failed to parse DOCX document: {str(e)}")
        return "\n".join(text_parts).strip()


def get_parser_for_extension(extension: str) -> Optional[DocumentParser]:
    """Factory helper to obtain the appropriate parser for a file extension."""
    ext = extension.lower()
    if ext == ".pdf":
        return PDFParser()
    elif ext == ".docx":
        return DOCXParser()
    return None
